"""MOP CNC + MultiExtracoes: worker unico, opt-in e sequencial.

IMPORTANTE:
  - Preflight nao acessa Supabase nem executa robo.
  - A execucao real exige --execute E ENABLE_UNIFIED_WORKER=SIM.
  - NUNCA executar simultaneamente ao worker CNC anterior.
  - SQLs de integracao/sql devem ter sido revisados/aplicados.
  - RUN_MULTI_PATH deve implementar o contrato documentado.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import socket
import subprocess
import sys
import tempfile
import time
from pathlib import Path

BASE = Path(__file__).resolve().parent
PRODUTOS = {
    "FGTS", "CNC", "CONSIGNADO", "ANTECIPAÇÃO DE BENEFÍCIO",
    "BMG CARD", "CARTAO BENEFICIO", "CONSIGNADO PRIVADO",
}
BUCKET = "mop-arquivos"
CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_XLSX_BYTES = 50 * 1024 * 1024


def env():
    from dotenv import load_dotenv
    load_dotenv(BASE / ".env")
    return {
        "url": os.getenv("SUPABASE_URL", ""),
        "secret": os.getenv("SUPABASE_SECRET_KEY", ""),
        "cnc": Path(os.getenv("RUN_MOP_PATH", "RUNTIME_NAO_CONFIGURADO")).expanduser(),
        "multi": Path(os.getenv("RUN_MULTI_PATH", "RUNTIME_NAO_CONFIGURADO")).expanduser(),
        "python": Path(os.getenv("PYTHON_EXE", sys.executable)).expanduser(),
        "poll": int(os.getenv("POLL_SECONDS", "10")),
        "timeout": int(os.getenv("JOB_TIMEOUT_SECONDS", "3600")),
        "worker_id": os.getenv("WORKER_ID", socket.gethostname() + "-mop-unificado"),
    }


def preflight(config):
    errors = []
    if not config["url"].startswith("https://") or ".supabase.co" not in config["url"]:
        errors.append("SUPABASE_URL nao configurada")
    if not config["secret"] or "COLE_AQUI" in config["secret"]:
        errors.append("SUPABASE_SECRET_KEY nao configurada")
    for field in ("cnc", "multi", "python"):
        if not config[field].is_file():
            errors.append(field.upper() + " nao encontrado: " + str(config[field]))
    if config["poll"] < 1 or config["timeout"] < 30:
        errors.append("POLL_SECONDS/JOB_TIMEOUT_SECONDS invalidos")
    if len(config["worker_id"].strip()) < 3:
        errors.append("WORKER_ID invalido")
    if errors:
        for msg in errors:
            print("[BLOQUEADO]", msg)
        return False
    print("[PREFLIGHT_OK] executores CNC e Multi encontrados; SEM CONEXAO OU EXECUCAO")
    return True


class SingleInstance:
    """Impede dois workers no mesmo computador; nao cobre outra maquina."""
    def __enter__(self):
        self.handle = (BASE / ".worker_unificado.lock").open("a+b")
        self.handle.seek(0, 2)
        if self.handle.tell() == 0:
            self.handle.write(b"0")
            self.handle.flush()
        self.handle.seek(0)
        try:
            if os.name == "nt":
                import msvcrt
                msvcrt.locking(self.handle.fileno(), msvcrt.LK_NBLCK, 1)
            else:
                import fcntl
                fcntl.flock(self.handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
        except OSError:
            self.handle.close()
            raise RuntimeError("Ja existe outro worker unificado nesta maquina")
        return self

    def __exit__(self, *_):
        self.handle.close()


def rpc(db, name, **kwargs):
    return db.rpc(name, kwargs).execute().data


def execute_process(command, timeout):
    return subprocess.run(command, capture_output=True, text=True,
                          encoding="utf-8", errors="replace",
                          timeout=timeout, check=False)


def process_cnc(db, config, job):
    job_id = int(job["id"])
    if not rpc(db, "mop_pedido_autorizado", p_id=job_id):
        rpc(db, "finish_mop", p_id=job_id, p_worker_id=config["worker_id"],
            p_success=False, p_error="ACESSO_LOJA_REVOGADO")
        return
    result = execute_process([
        str(config["python"]), str(config["cnc"]),
        "--franquia", str(job["franquia_codigo"]),
        "--whatsapp", str(job["whatsapp"])
    ], config["timeout"])
    err = None if result.returncode == 0 else (
        result.stderr.strip() or "Extracao CNC falhou"
    )[-3500:]
    rpc(db, "finish_mop", p_id=job_id, p_worker_id=config["worker_id"],
        p_success=result.returncode == 0, p_error=err)


def checked_multi_result(path, output_dir, order):
    """Nao confiar apenas no codigo 0: exigir XLSX e metadados validos."""
    from openpyxl import load_workbook
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("RESULTADO_JSON_INVALIDO")
    if data.get("pedido_id") != int(order["id"]):
        raise ValueError("PEDIDO_ID_INCOMPATIVEL")
    if data.get("produto") != order["produto"]:
        raise ValueError("PRODUTO_INCOMPATIVEL")
    if str(data.get("franquia_codigo")) != str(order["franquia_codigo"]):
        raise ValueError("FRANQUIA_INCOMPATIVEL")
    count = data.get("registros")
    if not isinstance(count, int) or isinstance(count, bool) or count < 1:
        raise ValueError("REGISTROS_INVALIDOS")
    name = data.get("arquivo")
    if not isinstance(name, str) or Path(name).name != name or not name.lower().endswith(".xlsx"):
        raise ValueError("NOME_XLSX_INVALIDO")
    file = (output_dir / name).resolve()
    if file.parent != output_dir.resolve() or not file.is_file():
        raise ValueError("XLSX_FORA_DO_DIRETORIO")
    if not (100 <= file.stat().st_size <= MAX_XLSX_BYTES):
        raise ValueError("TAMANHO_XLSX_INVALIDO")
    # Leitura estrutural: nao atesta qualidade dos dados extraidos.
    wb = load_workbook(file, read_only=True, data_only=True)
    try:
        if not wb.sheetnames:
            raise ValueError("XLSX_SEM_PLANILHA")
    finally:
        wb.close()
    return file, count


def process_multi(db, config, job):
    job_id = int(job["id"])
    if job["produto"] not in PRODUTOS:
        raise ValueError("PRODUTO_NAO_HOMOLOGADO")
    if not rpc(db, "extracao_mop_autorizada", p_id=job_id):
        rpc(db, "finish_extracoes_mop", p_id=job_id,
            p_worker_id=config["worker_id"], p_status="revisao_manual",
            p_registros=None, p_erro_publico="Acesso a franquia revogado")
        return
    with tempfile.TemporaryDirectory(prefix="mop_multi_") as temp:
        output = Path(temp).resolve()
        result_path = output / "resultado.json"
        result = execute_process([
            str(config["python"]), str(config["multi"]),
            "--pedido-id", str(job_id),
            "--franquia", str(job["franquia_codigo"]),
            "--produto", str(job["produto"]),
            "--diretorio-saida", str(output),
            "--resultado-json", str(result_path),
        ], config["timeout"])
        if result.returncode != 0 or not result_path.is_file():
            raise RuntimeError("EXECUTOR_MULTI_FALHOU")
        xlsx, records = checked_multi_result(result_path, output, job)
        content = xlsx.read_bytes()
        digest = hashlib.sha256(content).hexdigest()
        storage_path = "extracoes/{}/{}/relatorio.xlsx".format(job["user_id"], job_id)
        # Somente chave secreta server-side; nunca enviada para o navegador.
        db.storage.from_(BUCKET).upload(
            storage_path, content,
            file_options={"content-type": CONTENT_TYPE, "upsert": "false"},
        )
        db.table("arquivos_extracoes_mop").insert({
            "pedido_id": job_id, "nome": xlsx.name,
            "storage_path": storage_path, "sha256": digest,
            "tamanho_bytes": len(content),
        }).execute()
        rpc(db, "finish_extracoes_mop", p_id=job_id,
            p_worker_id=config["worker_id"], p_status="validado",
            p_registros=records, p_erro_publico=None)


def process(db, config, claim):
    origin = claim["origem"]
    job = claim["pedido"]
    if origin not in ("cnc", "multi") or not isinstance(job, dict):
        raise ValueError("CLAIM_INVALIDO")
    print("[INICIO] origem={} pedido={}".format(origin, job["id"]), flush=True)
    try:
        if origin == "cnc":
            process_cnc(db, config, job)
        else:
            process_multi(db, config, job)
    except Exception as exc:
        print("[FALHA] pedido={} classe={}".format(job["id"], type(exc).__name__),
              file=sys.stderr, flush=True)
        # Nao repetir automaticamente pedidos apos falha/timeout.
        try:
            if origin == "cnc":
                rpc(db, "finish_mop", p_id=int(job["id"]),
                    p_worker_id=config["worker_id"], p_success=False,
                    p_error="Falha de execucao; verificar log local")
            else:
                rpc(db, "finish_extracoes_mop", p_id=int(job["id"]),
                    p_worker_id=config["worker_id"], p_status="revisao_manual",
                    p_registros=None, p_erro_publico="Verificacao manual necessaria")
        except Exception:
            print("[ATENCAO] nao foi possivel atualizar status; requer reconciliacao",
                  file=sys.stderr)
        raise
    print("[FINAL] origem={} pedido={}".format(origin, job["id"]), flush=True)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--preflight", action="store_true")
    parser.add_argument("--execute", action="store_true")
    group = parser.add_mutually_exclusive_group()
    group.add_argument("--once", action="store_true")
    group.add_argument("--loop", action="store_true")
    args = parser.parse_args()
    cfg = env()
    if not preflight(cfg):
        return 2
    if not args.execute:
        print("[SEGURO] apenas preflight; nao houve claim ou extracao")
        return 0
    if not (args.once or args.loop) or os.getenv("ENABLE_UNIFIED_WORKER") != "SIM":
        print("[BLOQUEADO] requer --execute com --once/--loop e ENABLE_UNIFIED_WORKER=SIM")
        return 3
    from supabase import create_client
    db = create_client(cfg["url"], cfg["secret"])
    with SingleInstance():
        while True:
            claims = rpc(db, "claim_next_unificado_mop",
                         p_worker_id=cfg["worker_id"]) or []
            if claims:
                process(db, cfg, claims[0])
            if args.once:
                break
            if not claims:
                time.sleep(cfg["poll"])
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
